"""참석기업 리스트 엑셀의 마지막 열(P)에 매칭 국유특허를 기록한다."""
import json, sys
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment

src, out = sys.argv[1], sys.argv[2]
here = Path(__file__).parent
load = lambda n: json.loads((here / n).read_text(encoding="utf-8")) if (here / n).exists() else {}
selection, kipris, ais = load("selection.json"), load("kipris.json"), load("ais.json")
NUM = "①②③④⑤"

wb = openpyxl.load_workbook(src)
ws = wb.worksheets[0]
hdr_row = next(r for r in range(1, 20) if ws.cell(r, 3).value == "기업코드")
col = 16  # P열
ref = ws.cell(hdr_row, 13)  # '아이템' 헤더 서식을 그대로 사용
h = ws.cell(hdr_row, col, "매칭 국유특허(기술이전 추천)")
for a in ("font", "fill", "border", "alignment", "number_format"):
    setattr(h, a, copy(getattr(ref, a)))

for r in range(hdr_row + 1, ws.max_row + 1):
    code = ws.cell(r, 3).value
    apps = selection.get(code)
    if not apps:
        continue
    lines = []
    for i, a in enumerate(apps):
        k = kipris[a]
        no = f"등록 {k['regNo']}" if k["status"] == "등록" else f"공개 {k['appNo']}"
        score = (ais.get(a) or {}).get("score")
        lines.append(f"{NUM[i]} {k['title']} ({no})" + (f" [AIS 종합 {score}점]" if score is not None else ""))
    c = ws.cell(r, col, "\n".join(lines))
    body = ws.cell(r, 13)
    c.font, c.border = copy(body.font), copy(body.border)
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.column_dimensions["P"].width = 80
wb.save(out)

# 제안서 생성용 기업 정보 추출
companies = []
for r in range(hdr_row + 1, ws.max_row + 1):
    code = ws.cell(r, 3).value
    if code in selection:
        g = lambda c: ws.cell(r, c).value or ""
        companies.append({"code": code, "manager": g(4), "name": g(5), "cat": g(10),
                          "field": g(11), "sub": g(12), "item": g(13)})
(Path(__file__).parent / "companies.json").write_text(
    json.dumps(companies, ensure_ascii=False, indent=1), encoding="utf-8")
